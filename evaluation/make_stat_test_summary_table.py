"""Build the formatted statistical-test summary tables behind paper Table 6:
per dataset (Avg. over horizons) and per configuration (dataset x horizon),
ST-Mixer vs PatchTST / TimeMixer, with mean ± std, significance levels and
mean-difference 95% CIs from Welch's t-test on seed-level (3 runs) results.

Input (--error-bar): Excel workbook in the layout of `ERROR BAR.xlsx`
(sheet `Sheet1`, NO header row):
  - sections introduced by a model name in column 0 (STMIXER / PATCHTST /
    TIMEMIXER, mapped to ST-Mixer / PatchTST / TimeMixer);
  - within a section, a dataset name in column 0 starts a dataset block
    (LTMS is normalised to LMTS);
  - column 1: prediction horizon (96/192/336/720) or `Avg.`;
  - columns 2-5: MSE mean, MSE std, MAE mean, MAE std over the 3 seeds.

Output (--output-dir): `Statistical_Test_Summary_Table.xlsx` (styled sheets)
plus a `LaTeX/` subfolder with booktabs tables. Run from the repo root:
    python evaluation/make_stat_test_summary_table.py --error-bar <xlsx> --output-dir <dir>
"""

import pandas as pd
import numpy as np
from scipy import stats
import os

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

import argparse
import os

# ERROR BAR.xlsx 中每个Avg.行的mean/std来自3次运行（seed级），检验在seed层面进行，每组n=3
N_SEEDS = 3

MODEL_MAP = {'STMIXER': 'ST-Mixer', 'PATCHTST': 'PatchTST', 'TIMEMIXER': 'TimeMixer'}

# 行顺序：LMTS置于首位（与其后数据集之间用加粗分界线隔开以示区分），
# 其余数据集仍按论文主表顺序（表内LTMS统一写作LMTS）
DATASET_ORDER = ['LMTS', 'Traffic', 'Electricity', 'Weather', 'ETTh1', 'ETTh2', 'ETTm1', 'ETTm2']

def parse_avg(file_path):
    """解析ERROR BAR.xlsx中每个模型的Avg.行（seed级平均的mean与std）"""
    df = pd.read_excel(file_path, sheet_name='Sheet1', header=None)

    records = []
    current_model = None
    current_dataset = None

    for _, row in df.iterrows():
        c0 = row[0]
        if pd.notna(c0):
            c0s = str(c0).strip()
            if c0s.upper() in MODEL_MAP:
                current_model = MODEL_MAP[c0s.upper()]
                current_dataset = None
                continue
            # 表内LTMS统一写作LMTS
            current_dataset = 'LMTS' if c0s.upper() == 'LTMS' else c0s

        if current_model is None or pd.isna(row[1]):
            continue
        if str(row[1]).strip() != 'Avg.':
            continue

        records.append({
            'model': current_model, 'dataset': current_dataset,
            'MSE_mean': float(row[2]), 'MSE_std': float(row[3]),
            'MAE_mean': float(row[4]), 'MAE_std': float(row[5]),
        })

    return pd.DataFrame(records)

def parse_config(file_path):
    """解析ERROR BAR.xlsx中每个模型的逐配置行（数据集-步长，排除Avg.行）"""
    df = pd.read_excel(file_path, sheet_name='Sheet1', header=None)

    records = []
    current_model = None
    current_dataset = None

    for _, row in df.iterrows():
        c0 = row[0]
        if pd.notna(c0):
            c0s = str(c0).strip()
            if c0s.upper() in MODEL_MAP:
                current_model = MODEL_MAP[c0s.upper()]
                current_dataset = None
                continue
            # 表内LTMS统一写作LMTS
            current_dataset = 'LMTS' if c0s.upper() == 'LTMS' else c0s

        if current_model is None or pd.isna(row[1]):
            continue
        horizon = str(row[1]).strip()
        if horizon == 'Avg.':
            continue

        records.append({
            'model': current_model, 'dataset': current_dataset,
            'pred_len': int(float(horizon)),
            'MSE_mean': float(row[2]), 'MSE_std': float(row[3]),
            'MAE_mean': float(row[4]), 'MAE_std': float(row[5]),
        })

    return pd.DataFrame(records)

def welch_diff_ci(m1, s1, m2, s2, n=N_SEEDS):
    """均值差(m1-m2)、其95% CI半宽与双侧p值（Welch–Satterthwaite自由度）"""
    v1, v2 = s1 * s1 / n, s2 * s2 / n
    se = np.sqrt(v1 + v2)
    diff = m1 - m2
    if se == 0:
        return diff, np.nan, np.nan
    dfw = (v1 + v2) ** 2 / (v1 ** 2 / (n - 1) + v2 ** 2 / (n - 1))
    p = 2.0 * stats.t.sf(abs(diff / se), dfw)
    h = stats.t.ppf(0.975, dfw) * se
    return diff, h, p

def fmt_ci(diff, h, p):
    """Δ ± 95%CI半宽 格式；* / ** 分别标记95%/99%水平显著。
    只标记ST-Mixer显著更好（Δ<0）的情形：Δ>0但显著（ST-Mixer反而更差）不打星，
    避免把"显著变差"读成"显著更好"。"""
    if np.isnan(h):
        return '–'
    stars = ('**' if p < 0.01 else ('*' if p < 0.05 else '')) if diff < 0 else ''
    return f'{diff:+.3f} ± {h:.3f}{stars}'

def fmt(mean, std):
    return f'{mean:.3f} ± {std:.3f}'

def conf_tier(p, diff):
    """置信水平分层（TimeMixer Table 9风格）。
    只给出ST-Mixer显著更好（diff<0）时的水平；diff>=0（含显著更差）一律记为'–'。"""
    if diff >= 0:
        return '–'
    if p < 0.01:
        return '99%'
    if p < 0.05:
        return '95%'
    return '–'

def build_row(avg_df, dataset, m1, m2):
    """构造一个数据集行的内容：两个模型各MSE/MAE字符串 + MSE/MAE各自的置信水平（仅ST-Mixer更好时给出）"""
    r1 = avg_df[(avg_df['model'] == m1) & (avg_df['dataset'] == dataset)].iloc[0]
    r2 = avg_df[(avg_df['model'] == m2) & (avg_df['dataset'] == dataset)].iloc[0]

    d_mse, _, p_mse = welch_diff_ci(r1['MSE_mean'], r1['MSE_std'], r2['MSE_mean'], r2['MSE_std'])
    d_mae, _, p_mae = welch_diff_ci(r1['MAE_mean'], r1['MAE_std'], r2['MAE_mean'], r2['MAE_std'])

    return [
        fmt(r1['MSE_mean'], r1['MSE_std']), fmt(r1['MAE_mean'], r1['MAE_std']),
        fmt(r2['MSE_mean'], r2['MSE_std']), fmt(r2['MAE_mean'], r2['MAE_std']),
        conf_tier(p_mse, d_mse), conf_tier(p_mae, d_mae),
    ], d_mse, d_mae, p_mse, p_mae

def build_ci_row(avg_df, dataset, m1, m2):
    """构造一个数据集行的内容：两个模型各MSE/MAE字符串 + 均值差95%CI（±形式）"""
    r1 = avg_df[(avg_df['model'] == m1) & (avg_df['dataset'] == dataset)].iloc[0]
    r2 = avg_df[(avg_df['model'] == m2) & (avg_df['dataset'] == dataset)].iloc[0]

    d_mse, h_mse, p_mse = welch_diff_ci(r1['MSE_mean'], r1['MSE_std'], r2['MSE_mean'], r2['MSE_std'])
    d_mae, h_mae, p_mae = welch_diff_ci(r1['MAE_mean'], r1['MAE_std'], r2['MAE_mean'], r2['MAE_std'])

    return [
        fmt(r1['MSE_mean'], r1['MSE_std']), fmt(r1['MAE_mean'], r1['MAE_std']),
        fmt(r2['MSE_mean'], r2['MSE_std']), fmt(r2['MAE_mean'], r2['MAE_std']),
        fmt_ci(d_mse, h_mse, p_mse), fmt_ci(d_mae, h_mae, p_mae),
    ], d_mse, d_mae, p_mse, p_mae

def build_sheet(ws, m1, m2, avg_df, ci=False):
    """按TimeMixer Table 9样式构建带合并表头和三线表边框的sheet；ci=True时改输出均值差95%CI（±形式）"""
    thin = Side(style='thin', color='000000')
    medium = Side(style='medium', color='000000')

    font = Font(name='Times New Roman', size=12)
    font_bold = Font(name='Times New Roman', size=12, bold=True)
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(horizontal='left', vertical='center')

    # 两行表头（A列上下两行分开写，模型名与Significance横向合并）
    ws['A1'] = 'Model'
    ws.merge_cells('B1:C1'); ws['B1'] = m1
    ws.merge_cells('D1:E1'); ws['D1'] = f'{m2} (2023)' if m2 == 'PatchTST' else m2
    ws.merge_cells('F1:G1'); ws['F1'] = 'Mean ± half-width (95% CI)' if ci else 'Significance'
    ws['A2'] = 'Dataset'
    ws['B2'] = 'MSE'; ws['C2'] = 'MAE'; ws['D2'] = 'MSE'; ws['E2'] = 'MAE'
    ws['F2'] = 'MSE'; ws['G2'] = 'MAE'

    for row in ws['A1:G2']:
        for c in row:
            c.font = font
            c.alignment = center

    # 数据行
    r = 3
    pvalues = []
    row_fn = build_ci_row if ci else build_row
    for ds in DATASET_ORDER:
        values, d_mse, d_mae, p_mse, p_mae = row_fn(avg_df, ds, m1, m2)
        pvalues.append((ds, d_mse, d_mae, p_mse, p_mae))
        ws.cell(row=r, column=1, value=ds)
        for j, v in enumerate(values, start=2):
            ws.cell(row=r, column=j, value=v)
        for c in ws[r]:
            c.font = font
            c.alignment = left if c.column == 1 else center
        r += 1
    last_row = r - 1

    # 三线表边框：顶部粗线、表头下细线、底部粗线
    for row in ws['A1:G1']:
        for c in row:
            c.border = Border(top=medium)
    for row in ws['A2:G2']:
        for c in row:
            c.border = Border(bottom=thin)
    for row in ws[f'A{last_row}:G{last_row}']:
        for c in row:
            c.border = Border(bottom=medium)

    # 列宽
    w_ci = 18 if ci else 12
    widths = {'A': 13, 'B': 16, 'C': 16, 'D': 16, 'E': 16, 'F': w_ci, 'G': w_ci}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # 脚注
    if ci:
        note = (f"Note: mean ± std over 3 runs. Δ = {m1} − {m2} (negative Δ means {m1} is better), "
                "shown as mean ± half-width of its 95% CI (Welch's t-test, Welch–Satterthwaite df). "
                "* / **: ST-Mixer is significantly better at the 95% / 99% levels; "
                "significant differences in favour of the baseline are left unmarked.")
    else:
        note = ("Note: mean ± std over 3 runs. Values show the level at which ST-Mixer is significantly better "
                "(Welch's t-test on seed-level averaged results), reported separately for MSE and MAE; "
                "'–' = ST-Mixer is not significantly better at the 95% level.")
    note_row = last_row + 2
    ws.cell(row=note_row, column=1, value=note)
    ws.cell(row=note_row, column=1).font = Font(name='Times New Roman', size=10, italic=True)

    return pvalues

def build_config_sheet(ws, m1, m2, cfg_df, ci=False):
    """逐配置（数据集-步长）检验sheet，样式与Avg.表一致；ci=True时改输出均值差95%CI（±形式）"""
    thin = Side(style='thin', color='000000')
    medium = Side(style='medium', color='000000')

    font = Font(name='Times New Roman', size=12)
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(horizontal='left', vertical='center')

    # 两行表头：Model合并覆盖Dataset+Horizon两列
    ws.merge_cells('A1:B1'); ws['A1'] = 'Model'
    ws.merge_cells('C1:D1'); ws['C1'] = m1
    ws.merge_cells('E1:F1'); ws['E1'] = f'{m2} (2023)' if m2 == 'PatchTST' else m2
    ws.merge_cells('G1:H1'); ws['G1'] = 'Mean ± half-width (95% CI)' if ci else 'Significance'
    ws['A2'] = 'Dataset'; ws['B2'] = 'Horizon'
    ws['C2'] = 'MSE'; ws['D2'] = 'MAE'; ws['E2'] = 'MSE'; ws['F2'] = 'MAE'
    ws['G2'] = 'MSE'; ws['H2'] = 'MAE'

    for row in ws['A1:H2']:
        for c in row:
            c.font = font
            c.alignment = center

    # 数据行：8个数据集 × 4个步长
    horizons = [96, 192, 336, 720]
    r = 3
    pvalues = []
    for ds in DATASET_ORDER:
        for h in horizons:
            r1 = cfg_df[(cfg_df['model'] == m1) & (cfg_df['dataset'] == ds) & (cfg_df['pred_len'] == h)].iloc[0]
            r2 = cfg_df[(cfg_df['model'] == m2) & (cfg_df['dataset'] == ds) & (cfg_df['pred_len'] == h)].iloc[0]
            if ci:
                d_mse, h_mse, p_mse = welch_diff_ci(r1['MSE_mean'], r1['MSE_std'], r2['MSE_mean'], r2['MSE_std'])
                d_mae, h_mae, p_mae = welch_diff_ci(r1['MAE_mean'], r1['MAE_std'], r2['MAE_mean'], r2['MAE_std'])
                extra = [fmt_ci(d_mse, h_mse, p_mse), fmt_ci(d_mae, h_mae, p_mae)]
            else:
                d_mse, _, p_mse = welch_diff_ci(r1['MSE_mean'], r1['MSE_std'], r2['MSE_mean'], r2['MSE_std'])
                d_mae, _, p_mae = welch_diff_ci(r1['MAE_mean'], r1['MAE_std'], r2['MAE_mean'], r2['MAE_std'])
                extra = [conf_tier(p_mse, d_mse), conf_tier(p_mae, d_mae)]
            pvalues.append((ds, h, d_mse, d_mae, p_mse, p_mae))

            ws.cell(row=r, column=1, value=ds)
            ws.cell(row=r, column=2, value=h)
            values = [
                fmt(r1['MSE_mean'], r1['MSE_std']), fmt(r1['MAE_mean'], r1['MAE_std']),
                fmt(r2['MSE_mean'], r2['MSE_std']), fmt(r2['MAE_mean'], r2['MAE_std']),
            ] + extra
            for j, v in enumerate(values, start=3):
                ws.cell(row=r, column=j, value=v)
            for c in ws[r]:
                c.font = font
                c.alignment = left if c.column == 1 else center
            r += 1
    last_row = r - 1

    # 三线表边框：顶部粗线、表头下细线、底部粗线
    for row in ws['A1:H1']:
        for c in row:
            c.border = Border(top=medium)
    for row in ws['A2:H2']:
        for c in row:
            c.border = Border(bottom=thin)
    for row in ws[f'A{last_row}:H{last_row}']:
        for c in row:
            c.border = Border(bottom=medium)

    # 列宽
    w_ci = 18 if ci else 12
    widths = {'A': 13, 'B': 10, 'C': 16, 'D': 16, 'E': 16, 'F': 16, 'G': w_ci, 'H': w_ci}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # 脚注
    if ci:
        note = (f"Note: mean ± std over 3 runs. Δ = {m1} − {m2} (negative Δ means {m1} is better), "
                "shown as mean ± half-width of its 95% CI (Welch's t-test, Welch–Satterthwaite df). "
                "* / **: ST-Mixer is significantly better at the 95% / 99% levels; "
                "significant differences in favour of the baseline are left unmarked.")
    else:
        note = ("Note: mean ± std over 3 runs. Values show the level at which ST-Mixer is significantly better "
                "(Welch's t-test on per-configuration seed-level results), reported separately for MSE and MAE; "
                "'–' = ST-Mixer is not significantly better at the 95% level.")
    note_row = last_row + 2
    ws.cell(row=note_row, column=1, value=note)
    ws.cell(row=note_row, column=1).font = Font(name='Times New Roman', size=10, italic=True)

    return pvalues

def tex_decorate(v1, v2):
    """两模型数值比较：较优加粗、次优下划线（返回两个格式化后的LaTeX单元格）"""
    s1, s2 = f'{v1:.3f}', f'{v2:.3f}'
    if s1 == s2:
        return s1, s2
    if v1 < v2:
        return r'\textbf{%s}' % s1, r'\underline{%s}' % s2
    return r'\underline{%s}' % s1, r'\textbf{%s}' % s2

def tex_decorate_std(v1, s1, v2, s2):
    """两模型各以 mean ± std 呈现；因含±符号无法用siunitx的S列按小数点对齐，改用居中c列。
    较优加粗、次优下划线（与论文主表一致）"""
    c1, c2 = f'{v1:.3f} $\\pm$ {s1:.3f}', f'{v2:.3f} $\\pm$ {s2:.3f}'
    if f'{v1:.3f}' == f'{v2:.3f}':
        return c1, c2
    if v1 < v2:
        return r'\textbf{%s}' % c1, r'\underline{%s}' % c2
    return r'\underline{%s}' % c1, r'\textbf{%s}' % c2

def fmt_tex_ci(diff, h, p):
    """LaTeX的Δ±95%CI半宽单元格，*/**标记95%/99%显著（只标记ST-Mixer显著更好，即diff<0且显著）。
    正数用\\hphantom{-}补负号宽度，无星/单星用\\hphantom补齐到双星宽度，
    使所有单元格等宽：居中列下±与星标严格对齐且不悬出单元格（不与竖线相撞）"""
    if np.isnan(h):
        return '--'
    d = round(diff, 3)
    d = 0.0 if d == 0 else d  # 消除-0.000
    sign = '' if d < 0 else r'\hphantom{-}'
    if diff < 0 and p < 0.01:
        tail = '{}^{**}'
    elif diff < 0 and p < 0.05:
        tail = '{}^{*}' + r'\hphantom{{}^{*}}'
    else:
        tail = r'\hphantom{{}^{**}}'
    return '${}{:.3f} \\pm {:.3f}{}$'.format(sign, d, h, tail)

def fmt_tex_ci_interval(diff, h, p):
    """LaTeX的Δ的95%CI区间形式单元格：[lower, upper]，*/**标记95%/99%显著（只标记ST-Mixer显著更好，即diff<0且显著）。
    区间两端的正数用\\hphantom{-}补负号宽度，无星/单星用\\hphantom补齐到双星宽度，
    使所有单元格等宽：居中列下方括号[]与星标严格对齐，不被负号宽度打乱"""
    if np.isnan(h):
        return '--'
    lo = round(diff - h, 3)
    hi = round(diff + h, 3)
    lo = 0.0 if lo == 0 else lo  # 消除-0.000
    hi = 0.0 if hi == 0 else hi
    sign_lo = '' if lo < 0 else r'\hphantom{-}'
    sign_hi = '' if hi < 0 else r'\hphantom{-}'
    if diff < 0 and p < 0.01:
        tail = '{}^{**}'
    elif diff < 0 and p < 0.05:
        tail = '{}^{*}' + r'\hphantom{{}^{*}}'
    else:
        tail = r'\hphantom{{}^{**}}'
    return '$[{}{:.3f}, {}{:.3f}]{}$'.format(sign_lo, lo, sign_hi, hi, tail)

def export_perconfig_tex(cfg_df, avg_df, m1, m2, out_path, ci=False, ci_style='pm', model_std=False):
    """将逐配置sheet导出为参考主表样式（multirow数据集+Horizon列+三线表）的LaTeX表格。
    ci_style: 'pm' = Δ±半宽形式；'interval' = 只输出Δ的95%CI区间[lower, upper]
    model_std: True时两个模型列改为 mean ± std（数值列改用居中c列，并整体resizebox到栏宽）"""
    horizons = [96, 192, 336, 720]
    if not ci:
        tail_header = 'Significance'
    elif ci_style == 'interval':
        tail_header = r'$\Delta$ (95\% CI)'
    else:
        # 模型列已承载 mean ± std 后，末组必须显式写明Δ，避免被误读为"某模型自身的区间"
        tail_header = r'$\Delta \pm$ half-width (95\% CI)'

    # 表头/标题：所有导出表的两个模型列均呈现 mean ± std，caption 需一并说明
    head = r'\caption{Per-configuration comparison of ST-Mixer and ' + m2 + r' across prediction lengths. '
    # 星标只用于"ST-Mixer显著更好"；Δ>0但显著（ST-Mixer反而更差）不打星，故caption必须写明，
    # 否则读者会把"无星"误读为"无显著差异"，反而掩盖对ST-Mixer不利的显著结果
    star_note = (r' $^{*}$/$^{**}$: ST-Mixer is significantly better than ' + m2 +
                 r" at the 95\%/99\% levels (Welch's $t$-test); significant differences in favour of " +
                 m2 + r' are left unmarked.}')
    model_note = ((r'The ' + m1 + r' and ' + m2 +
                   r' column groups report mean $\pm$ std over 3 runs ($n=3$); ') if model_std else '')

    if ci and ci_style == 'interval':
        # 不写"intervals excluding zero indicate significance"——与星标说明重复
        # (∗的定义即95%区间不含0)，星标负责显著性、区间负责效应方向与大小
        caption = (head + model_note +
                   r'The last column group reports the 95\% confidence interval of $\Delta$ = ST-Mixer $-$ ' + m2 +
                   r" (Welch's $t$-test, $n=3$ runs); negative intervals mean ST-Mixer is better." + star_note)
        label = 'tab:%s_vs_%s_perconfig_ci_interval' % (m1.replace('-', '').lower(), m2.lower())
    elif ci:
        caption = (head + model_note +
                   r'The last column group reports the 95\% confidence interval ($\Delta \pm$ half-width) of '
                   r'$\Delta$ = ST-Mixer $-$ ' + m2 +
                   r" (Welch's $t$-test, $n=3$ runs), where negative values mean ST-Mixer is better." + star_note)
        label = 'tab:%s_vs_%s_perconfig_ci' % (m1.replace('-', '').lower(), m2.lower())
    else:
        caption = (head + model_note +
                   r'Significance columns report the level at which ST-Mixer is significantly better than ' + m2 +
                   r" (Welch's $t$-test over 3 runs); '--' denotes ST-Mixer is not significantly better at the "
                   r'95\% level.}')
        label = 'tab:%s_vs_%s_perconfig' % (m1.replace('-', '').lower(), m2.lower())

    def cfg_row(model, ds, h):
        return cfg_df[(cfg_df['model'] == model) & (cfg_df['dataset'] == ds) & (cfg_df['pred_len'] == h)].iloc[0]

    def avg_row(model, ds):
        return avg_df[(avg_df['model'] == model) & (avg_df['dataset'] == ds)].iloc[0]

    # 数值列用siunitx的S列：按小数点居中对齐，粗体/下划线不破坏对齐；
    # model_std 时数值含 ± 符号，S列无法按小数点对齐，改用居中 c 列（各值均为0.xxx三位小数，居中即自然对齐）
    s_col = 'S[table-format=0.3, detect-weight]'
    num_spec = 'cc|cc' if model_std else s_col * 2 + '|' + s_col * 2
    lines = [
        r'% 需要: \usepackage{booktabs, multirow, graphicx}' if model_std
        else r'% 需要: \usepackage{booktabs, multirow, siunitx}',
        r'\begin{table}[!t]',
        r'\centering',
        caption,
        r'\label{%s}' % label,
    ]
    if model_std:
        # 模型列加入 ±std 后自然宽度明显增大，统一缩到栏宽（与12列合并表同一策略）；
        # 注意 resizebox 只可包住 tabular，不可包住 \caption/\label（caption 在盒中排版会报错）
        # 仅Significance尾列（ci=False）用7pt：该表尾列很窄（99%/--），实测自然宽度仅422pt，
        # 被 \resizebox{\textwidth} 放大1.29倍后表高≈715pt，加上caption即超出页高726pt
        # （Float too large），底部被裁；加大列间距使自然宽度接近栏宽，放大倍数下降、表高落回页内。
        # CI表尾列是较宽的Δ数值（实测自然宽度519~530pt），只放大1.03~1.05倍、表高≈582pt，
        # 本就在页内，保持4pt以免列间距变宽后字号被进一步压小
        tabcolsep = '4pt' if ci else '7pt'
        if not ci:
            lines += [
                r'% 列间距由默认4pt加大到7pt：本表尾列很窄（99%/--），自然宽度仅422pt，',
                r'% 被 \resizebox{\textwidth} 放大1.29倍后表高会超出页高（Float too large），底部被裁；',
                r'% 加大列间距使自然宽度接近栏宽，放大倍数下降，表高回到页面内',
            ]
        lines += [r'\setlength{\tabcolsep}{%s}' % tabcolsep, r'\resizebox{\textwidth}{!}{%']
    lines += [
        r'\small',
        # 全高竖线分组：数据集 | 步长 | 模型1 | 模型2 | 置信水平/Δ
        r'\begin{tabular}{c|c|' + num_spec + '|cc}',
        r'\toprule',
        r'\multicolumn{2}{c|}{Models} & \multicolumn{2}{c|}{%s} & \multicolumn{2}{c|}{%s} & \multicolumn{2}{c}{%s} \\' % (m1, m2, tail_header),
        # 表内中间横线均用两端缩进的cmidrule（线宽取\lightrulewidth=0.05em，与原midrule一致），
        # 使其短于顶部/底部通栏粗线；显式指定(lr)与线宽，不依赖cmidrule默认值
        r'\cmidrule[\lightrulewidth](lr){1-8}',
        r'\multicolumn{2}{c|}{Metric} & \multicolumn{1}{c}{MSE} & \multicolumn{1}{c|}{MAE} & \multicolumn{1}{c}{MSE} & \multicolumn{1}{c|}{MAE} & \multicolumn{1}{c}{MSE} & \multicolumn{1}{c}{MAE} \\',
        r'\cmidrule[\lightrulewidth](lr){1-8}',
    ]

    # 尾部列格式：Δ±半宽 或 95%CI区间（Significance表不受影响）
    tail_fn = fmt_tex_ci_interval if (ci and ci_style == 'interval') else fmt_tex_ci

    for di, ds in enumerate(DATASET_ORDER):
        for h in horizons:
            r1, r2 = cfg_row(m1, ds, h), cfg_row(m2, ds, h)
            if model_std:
                mse1, mse2 = tex_decorate_std(r1['MSE_mean'], r1['MSE_std'], r2['MSE_mean'], r2['MSE_std'])
                mae1, mae2 = tex_decorate_std(r1['MAE_mean'], r1['MAE_std'], r2['MAE_mean'], r2['MAE_std'])
            else:
                mse1, mse2 = tex_decorate(r1['MSE_mean'], r2['MSE_mean'])
                mae1, mae2 = tex_decorate(r1['MAE_mean'], r2['MAE_mean'])
            d_mse, h_mse, p_mse = welch_diff_ci(r1['MSE_mean'], r1['MSE_std'], r2['MSE_mean'], r2['MSE_std'])
            d_mae, h_mae, p_mae = welch_diff_ci(r1['MAE_mean'], r1['MAE_std'], r2['MAE_mean'], r2['MAE_std'])
            tail = (tail_fn(d_mse, h_mse, p_mse), tail_fn(d_mae, h_mae, p_mae)) \
                if ci else tuple(conf_tier(p, d).replace('%', r'\%').replace('–', '--')
                                 for p, d in ((p_mse, d_mse), (p_mae, d_mae)))
            prefix = r'\multirow{5}{*}{%s}' % ds if h == 96 else ''
            lines.append(r'%s & %d & %s & %s & %s & %s & %s & %s \\' % (prefix, h, mse1, mae1, mse2, mae2, tail[0], tail[1]))

        # Avg.行上方的横线：不再切入左侧数据集名列，自数据集名列右侧起向右延伸至表格最右侧；
        # (lr)两端按\cmidrulekern缩进，覆盖Avg.标签列与全部MSE/MAE指标列
        lines.append(r'\cmidrule[\lightrulewidth](lr){2-8}')
        ra1, ra2 = avg_row(m1, ds), avg_row(m2, ds)
        if model_std:
            mse1, mse2 = tex_decorate_std(ra1['MSE_mean'], ra1['MSE_std'], ra2['MSE_mean'], ra2['MSE_std'])
            mae1, mae2 = tex_decorate_std(ra1['MAE_mean'], ra1['MAE_std'], ra2['MAE_mean'], ra2['MAE_std'])
        else:
            mse1, mse2 = tex_decorate(ra1['MSE_mean'], ra2['MSE_mean'])
            mae1, mae2 = tex_decorate(ra1['MAE_mean'], ra2['MAE_mean'])
        d_mse, h_mse, p_mse = welch_diff_ci(ra1['MSE_mean'], ra1['MSE_std'], ra2['MSE_mean'], ra2['MSE_std'])
        d_mae, h_mae, p_mae = welch_diff_ci(ra1['MAE_mean'], ra1['MAE_std'], ra2['MAE_mean'], ra2['MAE_std'])
        tail = (tail_fn(d_mse, h_mse, p_mse), tail_fn(d_mae, h_mae, p_mae)) \
            if ci else tuple(conf_tier(p, d).replace('%', r'\%').replace('–', '--')
                             for p, d in ((p_mse, d_mse), (p_mae, d_mae)))
        lines.append(r' & Avg. & %s & %s & %s & %s & %s & %s \\' % (mse1, mae1, mse2, mae2, tail[0], tail[1]))

        if di < len(DATASET_ORDER) - 1:
            # 数据集块之间的横线：通栏跨度但两端缩进，短于顶部/底部粗线；
            # LMTS块（首位数据集）与其后数据集的分界线用固定1.5pt加粗（约为\heavyrulewidth的0.08em近2倍，
            # 比表顶/底粗线更醒目；用pt而非em，保证各表字号不同时线宽一致），以示区分
            rule_w = r'1.5pt' if ds == 'LMTS' else r'\lightrulewidth'
            lines.append(r'\cmidrule[%s](lr){1-8}' % rule_w)

    lines += [r'\bottomrule', r'\end{tabular}}' if model_std else r'\end{tabular}', r'\end{table}']

    with open(out_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')
    return out_path

def rebuild_missing_avg(avg_df, cfg_df):
    """Avg行缺失时自动重建（当前文件PatchTST/TimeMixer的Avg行被误清空）：
    mean = 4个步长均值的平均（与文件中完好Avg行核验一致，数学上精确）；
    std  = 4个步长std的平均（与原文件PatchTST全部8个Avg std核验吻合的规律）。
    注意：ST-Mixer完好Avg行的std是真seed级数值，通常小于该近似——如需精确值
    建议从原始日志重新填入ERROR BAR.xlsx后再跑本脚本。"""
    fixed = []
    for i, r in avg_df.iterrows():
        if pd.notna(r['MSE_mean']) and pd.notna(r['MAE_mean']):
            continue
        sub = cfg_df[(cfg_df['model'] == r['model']) & (cfg_df['dataset'] == r['dataset'])]
        if len(sub) != 4:
            continue
        avg_df.loc[i, 'MSE_mean'] = sub['MSE_mean'].mean()
        avg_df.loc[i, 'MAE_mean'] = sub['MAE_mean'].mean()
        avg_df.loc[i, 'MSE_std'] = sub['MSE_std'].mean()
        avg_df.loc[i, 'MAE_std'] = sub['MAE_std'].mean()
        fixed.append((r['model'], r['dataset']))
    return fixed

def main():
    parser = argparse.ArgumentParser(
        description='Build formatted statistical-test summary tables (paper Table 6).',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__)
    parser.add_argument('--error-bar', default='results/error_bar.xlsx',
                        help='Path to the ERROR BAR workbook (layout: see module docstring).')
    parser.add_argument('--output-dir', default='results/statistical_tests/Welch_Test_PerConfig_MultiSeed',
                        help='Directory for the summary workbook and the LaTeX/ subfolder.')
    args = parser.parse_args()

    output_dir = args.output_dir
    os.makedirs(output_dir, exist_ok=True)

    avg_df = parse_avg(args.error_bar)
    cfg_df = parse_config(args.error_bar)
    fixed = rebuild_missing_avg(avg_df, cfg_df)
    if fixed:
        print('以下Avg.行在源文件中缺失，已自动重建（mean=步长均值平均, std=步长std平均）:')
        for m, ds in fixed:
            print(f'  {m} - {ds}')
    print('Avg.行解析结果:')
    print(avg_df.to_string(index=False))

    wb = Workbook()
    wb.remove(wb.active)

    pairs = [('ST-Mixer', 'PatchTST'), ('ST-Mixer', 'TimeMixer')]
    all_p = {}
    all_pc = {}
    all_ci = {}
    all_pcc = {}
    for m1, m2 in pairs:
        ws = wb.create_sheet(f'{m1} vs {m2}')
        all_p[(m1, m2)] = build_sheet(ws, m1, m2, avg_df)

        # 逐配置（数据集-步长）sheet
        ws_pc = wb.create_sheet(f'{m2} PerConfig')
        all_pc[(m1, m2)] = build_config_sheet(ws_pc, m1, m2, cfg_df)

        # 均值差95%CI（±形式）sheet：不改动原表，新增
        ws_ci = wb.create_sheet(f'{m1} vs {m2} CI')
        all_ci[(m1, m2)] = build_sheet(ws_ci, m1, m2, avg_df, ci=True)

        ws_pcc = wb.create_sheet(f'{m2} PerConfig CI')
        all_pcc[(m1, m2)] = build_config_sheet(ws_pcc, m1, m2, cfg_df, ci=True)

    excel_path = os.path.join(output_dir, 'Statistical_Test_Summary_Table.xlsx')
    wb.save(excel_path)
    print(f'\n汇总表已保存: {excel_path}')

    # LaTeX导出（参考主表样式：multirow数据集+粗体/下划线+三线表）
    latex_dir = os.path.join(output_dir, 'LaTeX')
    os.makedirs(latex_dir, exist_ok=True)
    tex_files = []
    for m1, m2 in pairs:
        tex_files.append(export_perconfig_tex(
            cfg_df, avg_df, m1, m2, os.path.join(latex_dir, f'ST-Mixer_vs_{m2}_PerConfig.tex'),
            ci=False, model_std=True))
        tex_files.append(export_perconfig_tex(
            cfg_df, avg_df, m1, m2, os.path.join(latex_dir, f'ST-Mixer_vs_{m2}_PerConfig_CI.tex'),
            ci=True, model_std=True))
        # 新增：Δ列替换为其95%置信区间[lower, upper]的版本
        tex_files.append(export_perconfig_tex(
            cfg_df, avg_df, m1, m2, os.path.join(latex_dir, f'ST-Mixer_vs_{m2}_PerConfig_CI_Interval.tex'),
            ci=True, ci_style='interval', model_std=True))
    print('LaTeX表格已导出:')
    for p in tex_files:
        print(f'  {p}')

    # 控制台打印（Markdown格式便于查看）
    for m1, m2 in pairs:
        print(f'\n### {m1} vs {m2}\n')
        print(f'| Dataset | {m1} MSE | {m1} MAE | {m2} MSE | {m2} MAE | Conf. (MSE) | Conf. (MAE) |')
        print('|---|---|---|---|---|---|---|')
        for ds, d_mse, d_mae, p_mse, p_mae in all_p[(m1, m2)]:
            values, _, _, _, _ = build_row(avg_df, ds, m1, m2)
            print(f'| {ds} | ' + ' | '.join(values) + ' |')
        print('\n各数据集p值 (MSE / MAE):')
        for ds, d_mse, d_mae, p_mse, p_mae in all_p[(m1, m2)]:
            print(f'  {ds:<12} MSE p={p_mse:.4f}  MAE p={p_mae:.4f}')

        print(f'\n### {m1} vs {m2} — 均值差95%CI（新增CI sheet）\n')
        print('| Dataset | ΔMSE (95% CI) | ΔMAE (95% CI) |')
        print('|---|---|---|')
        for ds, d_mse, d_mae, p_mse, p_mae in all_ci[(m1, m2)]:
            values, _, _, _, _ = build_ci_row(avg_df, ds, m1, m2)
            print(f'| {ds} | {values[4]} | {values[5]} |')

        print(f'\n--- {m1} vs {m2} 逐配置 ---\n')
        print('| Dataset | Horizon | Conf. (MSE) | Conf. (MAE) |')
        print('|---|---|---|---|')
        for ds, h, d_mse, d_mae, p_mse, p_mae in all_pc[(m1, m2)]:
            print(f'| {ds} | {h} | {conf_tier(p_mse, d_mse)} | {conf_tier(p_mae, d_mae)} |')

if __name__ == '__main__':
    main()
